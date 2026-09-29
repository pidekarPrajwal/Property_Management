from openpyxl import load_workbook

from properties.columns import PROPERTY_ID, XLSX_WRITABLE_FIELDS


class XlsxFormatError(Exception):
    """The workbook itself cannot be read or is missing the identity column."""


def _header(value):
    text = '' if value is None else str(value).strip().lower()
    return text.replace(' ', '_').replace('-', '_')


def read_xlsx(uploaded_file):
    """Return header names and data rows. This function does not touch the database."""
    name = getattr(uploaded_file, 'name', '') or ''
    if not name.lower().endswith('.xlsx'):
        raise XlsxFormatError('Upload an .xlsx file.')
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception as exc:
        raise XlsxFormatError('The file is not a readable .xlsx workbook.') from exc

    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            raise XlsxFormatError('The workbook has no header row.')
        headers = [_header(cell) for cell in header_row]
        if PROPERTY_ID not in headers:
            raise XlsxFormatError('The workbook must include a property_id column.')

        rows = []
        for offset, values in enumerate(iterator, start=2):
            if values is None or all(cell is None or str(cell).strip() == '' for cell in values):
                rows.append({'excel_row': offset, 'empty': True, 'values': {}})
                continue
            record = {}
            for index, header in enumerate(headers):
                if not header or index >= len(values):
                    continue
                record[header] = values[index]
            rows.append({'excel_row': offset, 'empty': False, 'values': record})
        recognized = [header for header in headers if header in XLSX_WRITABLE_FIELDS or header == PROPERTY_ID]
        ignored = [header for header in headers if header and header not in recognized]
        return {
            'headers': headers,
            'recognized_columns': recognized,
            'ignored_columns': ignored,
            'rows': rows,
        }
    finally:
        workbook.close()
